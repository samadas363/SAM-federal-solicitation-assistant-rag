import json
from pathlib import Path

import pymupdf  # PyMuPDF
from docx import Document
from openpyxl import load_workbook

RAW = Path("data/raw")
OUT = Path("data/processed")
OUT.mkdir(parents=True, exist_ok=True)
OUT_FILE = OUT / "documents.jsonl"

MIN_CHARS = 50  # pages with less text than this are probably scanned images


def parse_pdf(path):
    pages = []
    with pymupdf.open(path) as pdf:
        for i, page in enumerate(pdf, start=1):
            pages.append((i, page.get_text().strip()))
    return pages


def parse_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [(1, "\n".join(parts))]


def parse_xlsx(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    pages = []
    for i, sheet in enumerate(wb.worksheets, start=1):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        pages.append((i, f"[Sheet: {sheet.title}]\n" + "\n".join(rows)))
    return pages


PARSERS = {".pdf": parse_pdf, ".docx": parse_docx, ".xlsx": parse_xlsx}


def main():
    total_files = total_pages = 0
    flagged = []

    with open(OUT_FILE, "w", encoding="utf-8") as out:
        for folder in sorted(p for p in RAW.iterdir() if p.is_dir()):
            meta_path = folder / "metadata.json"
            meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}
            title = meta.get("title", "")
            print(f"{folder.name} - {title[:60]}")

            for file in sorted(folder.iterdir()):
                parser = PARSERS.get(file.suffix.lower())
                if not parser:
                    continue
                try:
                    pages = parser(file)
                except Exception as e:
                    print(f"  FAILED {file.name}: {e}")
                    continue

                total_files += 1
                for page_no, text in pages:
                    if len(text) < MIN_CHARS:
                        flagged.append((file.name, page_no))
                        continue
                    record = {
                        "notice_id": folder.name,
                        "title": title,
                        "agency": meta.get("fullParentPathName", ""),
                        "naics": meta.get("naicsCode", ""),
                        "due_date": meta.get("responseDeadLine", ""),
                        "file": file.name,
                        "page": page_no,
                        "text": text,
                    }
                    out.write(json.dumps(record, ensure_ascii=False) + "\n")
                    total_pages += 1

    print(f"\nParsed {total_files} files, saved {total_pages} pages to {OUT_FILE}")
    print(f"{len(flagged)} pages had little or no text (possibly scanned)")
    for name, page in flagged[:10]:
        print(f"  {name} page {page}")


if __name__ == "__main__":
    main()