
import io
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook


def read_pdf(data: bytes):
    """Return (page_number, text) pairs from an uploaded PDF."""
    pages = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text().strip()
            if text:
                pages.append((page_number, text))
    return pages


def read_docx(data: bytes):
    """Return the text of an uploaded DOCX file."""
    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    text = "\n".join(part for part in parts if part)
    return [(1, text)] if text else []


def read_xlsx(data: bytes):
    """Return (sheet_number, text) pairs from an uploaded XLSX file."""
    pages = []
    wb = load_workbook(
        io.BytesIO(data),
        data_only=True,
        read_only=True,
    )

    try:
        for sheet_number, sheet in enumerate(wb.worksheets, start=1):
            rows = []

            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    rows.append(" | ".join(cells))

            if rows:
                pages.append(
                    (sheet_number, f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
                )
    finally:
        wb.close()

    return pages


def read_uploaded_file(filename: str, data: bytes):
    """Read a supported uploaded file and return page/sheet text."""
    suffix = Path(filename).suffix.lower()

    readers = {
        ".pdf": read_pdf,
        ".docx": read_docx,
        ".xlsx": read_xlsx,
    }

    reader = readers.get(suffix)
    if reader is None:
        raise ValueError("Supported formats are PDF, DOCX, and XLSX.")

    if not data:
        raise ValueError("The uploaded file is empty.")

    return reader(data)
